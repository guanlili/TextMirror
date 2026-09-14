"""修订文档生成测试：docx 逐段替换+红色标记、txt 全量替换。

修订件是用户直接下载的产物，坏了一般到用户手里才发现——此前零测试。
"""
import os
import tempfile

from docx import Document

from app.services.document import generate_corrected_docx, generate_corrected_txt

ISSUES = [
    {"original": "帐号", "suggestion": "账号"},
    {"original": "权力和义务", "suggestion": "权利和义务"},
    # 无建议/原文建议相同：不应进入替换
    {"original": "孤词", "suggestion": ""},
    {"original": "同词", "suggestion": "同词"},
]


def _make_doc(paragraphs: list) -> str:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    path = os.path.join(tempfile.mkdtemp(prefix="tm-corrected-"), "origin.docx")
    doc.save(path)
    return path


def test_corrected_docx_replaces_and_marks_red():
    src = _make_doc(["帐号管理规范", "无修改段落", "享有权力和义务"])
    out = src.replace("origin.docx", "corrected.docx")
    generate_corrected_docx(src, ISSUES, out)

    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "账号管理规范"          # 已替换
    assert texts[1] == "无修改段落"            # 未动
    assert texts[2] == "享有权利和义务"        # 搭配替换

    # 修改段红色标记，未修改段无红色
    assert doc.paragraphs[0].runs[0].font.color.rgb is not None
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "FF0000"
    assert doc.paragraphs[1].runs[0].font.color.rgb is None


def test_corrected_docx_no_issues_returns_original_content():
    src = _make_doc(["完全干净的文本"])
    out = src.replace("origin.docx", "corrected.docx")
    generate_corrected_docx(src, [], out)
    assert Document(out).paragraphs[0].text == "完全干净的文本"


def test_corrected_txt_replaces_all_occurrences():
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, "corrected.txt")
        generate_corrected_txt("帐号错了，帐号又错了", ISSUES, out)
        with open(out, encoding="utf-8") as f:
            assert f.read() == "账号错了，账号又错了"  # replace 全量替换
