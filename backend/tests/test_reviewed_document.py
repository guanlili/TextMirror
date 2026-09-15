"""局部 DOCX 补丁与格式保留；跨段/特殊内容明确拒绝而非生成错误文档。"""
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.services.document import _extract_docx, generate_reviewed_docx


def patch(source, original, replacement, start=None):
    start = source.index(original) if start is None else start
    return {"start": start, "end": start + len(original), "original": original, "replacement": replacement}


def test_cross_run_second_occurrence_emoji_and_styles(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(15)
    p.add_run("  😀错词，").bold = True
    p.add_run("第2处错").italic = True
    colored = p.add_run("词")
    colored.font.color.rgb = RGBColor.from_string("123456")
    untouched = p.add_run("尾巴  ")
    untouched.font.size = Pt(18)
    doc.add_paragraph("   ")
    doc.add_paragraph("后段")
    original, out = tmp_path / "original.docx", tmp_path / "out.docx"
    doc.save(original)
    source = _extract_docx(str(original))
    rpr = [run._r.rPr.xml for run in p.runs]
    ppr = p._p.pPr.xml
    untouched_xml = untouched._r.xml
    start = source.rindex("错词")
    generate_reviewed_docx(str(original), source, [patch(source, "错词", "正确词", start)], str(out))
    result = Document(out)
    assert _extract_docx(str(out)) == "😀错词，第2处正确词尾巴\n后段"
    assert result.paragraphs[0].text.startswith("  😀错词，")
    assert [run._r.rPr.xml for run in result.paragraphs[0].runs] == rpr
    assert result.paragraphs[0]._p.pPr.xml == ppr
    assert result.paragraphs[0].runs[-1]._r.xml == untouched_xml
    assert result.paragraphs[1].text == "   "


def test_multiple_patches_same_run_and_deletion(tmp_path):
    doc = Document()
    doc.add_paragraph("😀敏感，甲错乙错尾")
    original, out = tmp_path / "source.docx", tmp_path / "out.docx"
    doc.save(original)
    source = _extract_docx(str(original))
    patches = [patch(source, "敏感，", ""), patch(source, "错", "正确"),
               patch(source, "错", "正", source.rindex("错"))]
    generate_reviewed_docx(str(original), source, patches, str(out))
    assert _extract_docx(str(out)) == "😀甲正确乙正尾"


def test_unaccepted_export_byte_identical_including_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("完全原样")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格不在审校提取范围内"
    original, out = tmp_path / "original.docx", tmp_path / "out.docx"
    doc.save(original)
    generate_reviewed_docx(str(original), "完全原样", [], str(out))
    assert original.read_bytes() == out.read_bytes()


def test_hyperlink_and_run_properties_preserved(tmp_path):
    doc = Document()
    p = doc.add_paragraph("前错")
    hyperlink = OxmlElement("w:hyperlink")
    relation = p.part.relate_to("https://example.com/", RT.HYPERLINK, is_external=True)
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:b"))
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = "链接错词"
    run.append(node)
    hyperlink.append(run)
    p._p.append(hyperlink)
    p.add_run("尾")
    original, out = tmp_path / "original.docx", tmp_path / "out.docx"
    doc.save(original)
    source = _extract_docx(str(original))
    generate_reviewed_docx(str(original), source, [patch(source, "前错", "前正")], str(out))
    result = Document(out)
    link = result.paragraphs[0]._p.xpath("./w:hyperlink")[0]
    assert link.xml == hyperlink.xml
    assert result.part.rels[link.get(qn("r:id"))].target_ref == "https://example.com/"
    # A hit inside a hyperlink keeps its relationship and bold style, too.
    generate_reviewed_docx(str(original), source, [patch(source, "错词", "正词")], str(out))
    result = Document(out)
    link = result.paragraphs[0]._p.xpath("./w:hyperlink")[0]
    assert link.get(qn("r:id")) == relation and link.xpath(".//w:b")
    assert _extract_docx(str(out)) == source.replace("错词", "正词")


@pytest.mark.parametrize("paragraphs,original,replacement", [
    (["第一错", "第二段"], "错\n第", "正"),
    (["甲\n乙"], "甲\n乙", "丙"),
    (["甲乙"], "乙", "乙\n丙"),
    (["甲", "乙"], "甲", ""),  # 删除整个非空段落会使提取后的换行消失
    (["甲乙"], "甲", " "),  # trim 将导致导出正文与成稿不同
])
def test_unsafe_semantics_explicitly_rejected(tmp_path, paragraphs, original, replacement):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    origin, out = tmp_path / "original.docx", tmp_path / "out.docx"
    doc.save(origin)
    source = _extract_docx(str(origin))
    with pytest.raises(ValueError, match="TXT"):
        generate_reviewed_docx(str(origin), source, [patch(source, original, replacement)], str(out))
    assert not out.exists()


def test_changed_source_rejected_even_without_patches(tmp_path):
    doc = Document()
    doc.add_paragraph("源文件另一个版本")
    origin, out = tmp_path / "source.docx", tmp_path / "out.docx"
    doc.save(origin)
    with pytest.raises(ValueError, match="不一致"):
        generate_reviewed_docx(str(origin), "旧原文", [], str(out))


def append_field_marker(paragraph, field_type):
    marker = OxmlElement("w:fldChar")
    marker.set(qn("w:fldCharType"), field_type)
    paragraph.add_run()._r.append(marker)


@pytest.mark.parametrize("cached,original", [("错词", "错词"), ("错词", "前错词尾"), ("", "前尾")])
def test_dynamic_field_result_and_boundaries_rejected(tmp_path, cached, original):
    doc = Document()
    p = doc.add_paragraph("前")
    append_field_marker(p, "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = ' QUOTE "错词" '
    p.add_run()._r.append(instruction)
    append_field_marker(p, "separate")
    p.add_run(cached)
    append_field_marker(p, "end")
    p.add_run("尾")
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    origin, out = tmp_path / "field.docx", tmp_path / "out.docx"
    doc.save(origin)
    source = _extract_docx(str(origin))
    with pytest.raises(ValueError, match="TXT"):
        generate_reviewed_docx(str(origin), source, [patch(source, original, "正词")], str(out))
    assert not out.exists()


def test_field_spanning_paragraphs_protects_nested_result(tmp_path):
    doc = Document()
    p = doc.add_paragraph("前")
    append_field_marker(p, "begin")
    append_field_marker(p, "separate")
    p = doc.add_paragraph()
    append_field_marker(p, "begin")
    append_field_marker(p, "separate")
    p.add_run("内层")
    append_field_marker(p, "end")
    p.add_run("错词")
    append_field_marker(p, "end")
    origin, out = tmp_path / "field.docx", tmp_path / "out.docx"
    doc.save(origin)
    source = _extract_docx(str(origin))
    with pytest.raises(ValueError, match="TXT"):
        generate_reviewed_docx(str(origin), source, [patch(source, "错词", "正词")], str(out))
    assert not out.exists()


def test_plain_text_outside_field_can_be_edited(tmp_path):
    doc = Document()
    p = doc.add_paragraph("前错")
    append_field_marker(p, "begin")
    append_field_marker(p, "separate")
    p.add_run("域结果")
    append_field_marker(p, "end")
    p.add_run("尾错")
    field_xml = [run._r.xml for run in p.runs[1:-1]]
    origin, out = tmp_path / "field.docx", tmp_path / "out.docx"
    doc.save(origin)
    source = _extract_docx(str(origin))
    generate_reviewed_docx(str(origin), source, [patch(source, "前错", "前正"), patch(source, "尾错", "尾正")], str(out))
    result = Document(out)
    assert result.paragraphs[0].text == "前正域结果尾正"
    assert [run._r.xml for run in result.paragraphs[0].runs[1:-1]] == field_xml


def test_patch_cannot_span_unextracted_simple_field(tmp_path):
    doc = Document()
    p = doc.add_paragraph("前")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), ' QUOTE "域结果" ')
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "域结果"
    run.append(text)
    field.append(run)
    p._p.append(field)
    p.add_run("尾")
    origin, out = tmp_path / "field.docx", tmp_path / "out.docx"
    doc.save(origin)
    source = _extract_docx(str(origin))
    with pytest.raises(ValueError, match="TXT"):
        generate_reviewed_docx(str(origin), source, [patch(source, source, "正词")], str(out))
    assert not out.exists()
