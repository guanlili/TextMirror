"""审阅 API、CAS、版本与私有导出；SQLite/fakeredis，不调用外部模型。"""
import copy
import io
import os
import uuid
from unittest.mock import AsyncMock
from urllib.parse import unquote

import pytest
from docx import Document
from fastapi import HTTPException
from sqlalchemy import func, select

from app.core import redis as redis_module
from app.core.database import async_session_factory
from app.core.dependencies import get_current_user, get_current_user_optional
from app.main import app
from app.models.proofread import ProofreadRecord
from app.models.proofread_task import ProofreadTask
from app.models.role import Role
from app.models.uploaded_document import UploadedDocument
from app.models.user import User
from app.schemas.review import ReviewWriteRequest
from app.services.review import load_review_record, save_review

SOURCE = "甲😀错词，乙错词。尾"


def issue(**changes):
    value = {"original": "错词", "suggestion": "正词", "type": "typo", "severity": "warning",
             "start": 6, "end": 8, "_accepted": True}
    value.update(changes)
    return value


@pytest.fixture
async def owner(client):
    async with async_session_factory() as db:
        role = Role(name="审阅测试", code=uuid.uuid4().hex)
        db.add(role)
        await db.flush()
        user = User(employee_id=uuid.uuid4().hex, username="审阅用户", password_hash="unused", role_id=role.id)
        db.add(user)
        await db.commit()
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_current_user_optional] = lambda: user
    yield user
    app.dependency_overrides.pop(get_current_user, None)
    app.dependency_overrides.pop(get_current_user_optional, None)


async def record_for(owner, **kwargs):
    defaults = dict(user_id=owner.id, type="text", original_text=SOURCE,
                    result={"issues": [issue(start=None, end=None, _accepted=False)], "depth": "deep", "config_id": 7})
    defaults.update(kwargs)
    async with async_session_factory() as db:
        record = ProofreadRecord(**defaults)
        db.add(record)
        await db.commit()
        return record


async def upload_for(owner, tmp_path, **kwargs):
    filename = "中文原稿.txt"
    path = tmp_path / filename
    path.write_text(SOURCE, encoding="utf-8")
    values = dict(file_id=str(uuid.uuid4()), filename=filename, file_ext=".txt", file_path=str(path),
                  extracted_text=SOURCE, owner_kind="user" if owner else "guest", user_id=owner.id if owner else None)
    values.update(kwargs)
    async with async_session_factory() as db:
        document = UploadedDocument(**values)
        db.add(document)
        await db.commit()
        return document


async def test_review_roundtrip_exact_unicode_versions_and_no_quota(client, owner):
    record = await record_for(owner, modified_text="不可沿用的自动成稿")
    url = f"/api/v1/history/{record.id}"
    first = (await client.get(url + "/review")).json()
    assert first["revision"] == 0 and first["modified_text"] == SOURCE
    assert first["issues"][0]["start"] is None
    assert first["coverage"] is None and first["depth"] == "deep" and first["config_id"] == 7
    assert first["source_file_id"] is None
    response = await client.put(url + "/review", json={"revision": 0, "issues": [issue()]})
    assert response.status_code == 200, response.text
    assert response.json()["modified_text"] == "甲😀错词，乙正词。尾"
    assert response.json()["issues"][0]["_accepted"] is True
    assert "accepted" not in response.json()["issues"][0]
    saved = await client.post(url + "/versions", json={"revision": 1, "issues": [issue()]})
    assert saved.status_code == 200, saved.text
    version = saved.json()["versions"][0]
    assert version["label"] and version["created_at"] and version["id"]
    reverted = await client.put(url + "/review", json={"revision": 2, "issues": [issue(_accepted=False)]})
    assert reverted.json()["modified_text"] == SOURCE
    assert reverted.json()["versions"][0] == version
    assert (await client.get(url)).json()["modified_text"] == SOURCE
    exported = await client.post(url + "/export", json={"revision": 3, "version_id": version["id"], "format": "txt"})
    assert exported.text == version["modified_text"]
    assert "file_path" not in reverted.text
    async with async_session_factory() as db:
        count = await db.scalar(select(func.count()).select_from(ProofreadRecord).where(ProofreadRecord.user_id == owner.id))
        original = await db.get(ProofreadRecord, record.id)
        assert count == 1 and original.original_text == SOURCE and original.result == record.result
    assert not await redis_module.redis_client.keys("textmirror:user_daily:*")


@pytest.mark.parametrize("changes", [
    {"start": 2, "end": 4, "original": "不符"}, {"start": 6, "end": 999},
    {"start": -1, "end": -1}, {"start": 6, "end": None}, {"start": None, "end": None},
    {"_ignored": True}, {"suggestion": ""}, {"start": 6.0},
    {"_patch": {"start": 0, "end": 1, "replacement": "伪造"}},
])
async def test_invalid_issue_rejected(client, owner, changes):
    record = await record_for(owner)
    result = await client.put(f"/api/v1/history/{record.id}/review", json={"revision": 0, "issues": [issue(**changes)]})
    assert result.status_code == 422, result.text


async def test_pending_unlocated_and_sensitive_single_deletion(client, owner):
    record = await record_for(owner)
    response = await client.put(f"/api/v1/history/{record.id}/review", json={"revision": 0, "issues": [
        issue(type="sensitive", suggestion=""),
        issue(start=None, end=None, original="未定位", _accepted=False),
    ]})
    assert response.status_code == 200, response.text
    assert response.json()["modified_text"] == "甲😀错词，乙尾"


async def test_shared_patch_from_multiple_issue_types_applied_once(client, owner):
    record = await record_for(owner)
    response = await client.put(f"/api/v1/history/{record.id}/review", json={"revision": 0, "issues": [
        issue(), issue(type="grammar"),
    ]})
    assert response.status_code == 200, response.text
    assert response.json()["modified_text"] == "甲😀错词，乙正词。尾"
    assert len(response.json()["issues"]) == 2


async def test_unlocated_issues_do_not_share_ignore_decisions(client, owner):
    record = await record_for(owner)
    response = await client.put(f"/api/v1/history/{record.id}/review", json={"revision": 0, "issues": [
        issue(start=None, end=None, _accepted=False, _ignored=True),
        issue(start=None, end=None, type="grammar", _accepted=False, _ignored=False),
    ]})
    assert response.status_code == 200, response.text
    assert response.json()["modified_text"] == SOURCE
    assert [item["_ignored"] for item in response.json()["issues"]] == [True, False]


async def test_overlap_including_swallowed_punctuation_rejected(client, owner):
    record = await record_for(owner)
    response = await client.put(f"/api/v1/history/{record.id}/review", json={"revision": 0, "issues": [
        issue(type="sensitive", suggestion=""), issue(start=8, end=9, original="。", suggestion="！"),
    ]})
    assert response.status_code == 409


async def test_revision_conflict_and_version_limit(client, owner):
    record = await record_for(owner)
    url = f"/api/v1/history/{record.id}"
    for revision in range(20):
        response = await client.post(url + "/versions", json={"revision": revision, "issues": []})
        assert response.status_code == 200, response.text
    full = response.json()
    response = await client.post(url + "/versions", json={"revision": 20, "issues": []})
    assert response.status_code == 422 and "20" in response.text
    assert (await client.get(url + "/review")).json() == full
    assert (await client.put(url + "/review", json={"revision": 0, "issues": []})).status_code == 409
    assert (await client.post(url + "/export", json={"revision": 0, "format": "txt"})).status_code == 409
    assert (await client.post(url + "/export", json={"revision": 20, "format": "txt", "version_id": "missing"})).status_code == 404
    # Version ceiling does not prevent updating a draft.
    assert (await client.put(url + "/review", json={"revision": 20, "issues": []})).status_code == 200


async def test_cas_prevents_stale_loaded_record_overwrite(client, owner):
    record = await record_for(owner)
    async with async_session_factory() as first, async_session_factory() as second:
        stale = await load_review_record(first, record.id, owner.id)
        # End SQLite read transaction but retain stale identity-map object, simulating a delayed writer.
        await first.commit()
        latest = await load_review_record(second, record.id, owner.id)
        await save_review(second, latest, owner.id, ReviewWriteRequest(revision=0, issues=[issue()]))
        await second.commit()
        with pytest.raises(HTTPException) as failure:
            await save_review(first, stale, owner.id, ReviewWriteRequest(revision=0, issues=[]))
        assert failure.value.status_code == 409
        await first.rollback()
    assert (await client.get(f"/api/v1/history/{record.id}/review")).json()["modified_text"] != SOURCE


async def test_ownership_record_types_and_anonymous(client, owner):
    foreign = await record_for(owner, user_id=None)
    polish = await record_for(owner, type="polish")
    for record in (foreign, polish):
        root = f"/api/v1/history/{record.id}"
        assert (await client.get(root + "/review")).status_code == 404
        assert (await client.put(root + "/review", json={"revision": 0, "issues": []})).status_code == 404
        assert (await client.post(root + "/versions", json={"revision": 0, "issues": []})).status_code == 404
        assert (await client.post(root + "/export", json={"revision": 0, "format": "txt"})).status_code == 404
    app.dependency_overrides.pop(get_current_user)
    assert (await client.get(f"/api/v1/history/{foreign.id}/review")).status_code in (401, 403)


async def test_coverage_and_payload_limits(client, owner):
    record = await record_for(owner, original_text="甲" * 100000)
    url = f"/api/v1/history/{record.id}/review"
    coverage = {"status": "partial", "total_chunks": 2, "completed_chunks": 1,
                "failed_chunks": [{"chunk_index": 1, "start": 50000, "end": 100000, "text": "甲" * 50000}]}
    pending = issue(start=None, end=None, _accepted=False)
    response = await client.put(url, json={"revision": 0, "issues": [pending] * 3000, "coverage": coverage})
    assert response.status_code == 200, response.text
    assert len(response.json()["original_text"]) == 100000
    invalid = copy.deepcopy(coverage)
    invalid["failed_chunks"][0]["text"] = "伪造"
    assert (await client.put(url, json={"revision": 1, "issues": [], "coverage": invalid})).status_code == 422
    invalid = {**coverage, "status": "complete"}
    assert (await client.put(url, json={"revision": 1, "issues": [], "coverage": invalid})).status_code == 422
    assert (await client.put(url, json={"revision": 1, "issues": [pending] * 5001})).status_code == 422
    oversized = [{**pending, "explanation": "甲" * 20000}] * 150
    assert (await client.put(url, json={"revision": 1, "issues": oversized})).status_code == 422
    assert (await client.put(url, json={"revision": 1, "issues": [], "modified_text": "伪造成稿"})).status_code == 422


@pytest.mark.parametrize("missing", ["legacy", "disk", "deleted", "foreign", "apikey"])
async def test_history_missing_source_docx_410_txt_works(client, owner, tmp_path, missing):
    document = await upload_for(owner, tmp_path)
    if missing == "disk":
        os.unlink(document.file_path)
    async with async_session_factory() as db:
        row = await db.get(UploadedDocument, document.id)
        if missing == "deleted":
            row.status = "deleted"
        if missing == "foreign":
            row.user_id = None
        if missing == "apikey":
            row.owner_kind = "apikey"
        await db.commit()
    record = await record_for(owner, type="document", source_filename=document.filename,
                              source_file_id=None if missing == "legacy" else document.file_id)
    url = f"/api/v1/history/{record.id}"
    await client.put(url + "/review", json={"revision": 0, "issues": []})
    assert (await client.post(url + "/export", json={"revision": 1, "format": "docx"})).status_code == 410
    result = await client.post(url + "/export", json={"revision": 1, "format": "txt"})
    assert result.status_code == 200 and result.text == SOURCE
    assert "中文原稿" in unquote(result.headers["content-disposition"])


async def test_text_docx_and_unsaved_export(client, owner):
    record = await record_for(owner)
    url = f"/api/v1/history/{record.id}"
    assert (await client.post(url + "/export", json={"revision": 0, "format": "txt"})).status_code == 422
    await client.put(url + "/review", json={"revision": 0, "issues": [issue()]})
    result = await client.post(url + "/export", json={"revision": 1, "format": "docx"})
    assert result.status_code == 200
    assert Document(io.BytesIO(result.content)).paragraphs[0].text == "甲😀错词，乙正词。尾"


async def test_guest_export_policy_ownership_and_no_model(client, tmp_path, monkeypatch):
    from app.api.v1 import document as document_api
    enabled = AsyncMock()
    monkeypatch.setattr(document_api, "reject_guest_if_disabled", enabled)
    forbidden_model = AsyncMock(side_effect=AssertionError("must not call LLM"))
    monkeypatch.setattr(document_api, "proofread_text", forbidden_model)
    document = await upload_for(None, tmp_path)
    url = f"/api/v1/document/{document.file_id}/export"
    response = await client.post(url, json={"issues": [issue()], "format": "txt"})
    assert response.status_code == 200 and response.text == "甲😀错词，乙正词。尾"
    enabled.assert_awaited_once()
    forbidden_model.assert_not_called()
    await redis_module.redis_client.hset("system:policy:guest", mapping={"allow_upload": "0"})
    assert (await client.post(url, json={"issues": [], "format": "txt"})).status_code == 403
    await redis_module.redis_client.hset("system:policy:guest", mapping={"allow_upload": "1"})
    for kind in ("apikey", "legacy", "user"):
        async with async_session_factory() as db:
            row = await db.get(UploadedDocument, document.id)
            row.owner_kind = kind
            await db.commit()
        assert (await client.post(url, json={"issues": [], "format": "txt"})).status_code == 404
    monkeypatch.setattr(document_api, "reject_guest_if_disabled", AsyncMock(side_effect=HTTPException(403, "disabled")))
    assert (await client.post(url, json={"issues": [], "format": "txt"})).status_code == 403


async def test_export_temporary_file_cleanup(client, owner, monkeypatch):
    import app.services.review as service
    paths = []
    original = service._write_export

    def capture(*args):
        paths.append(args[3])
        return original(*args)

    monkeypatch.setattr(service, "_write_export", capture)
    record = await record_for(owner)
    url = f"/api/v1/history/{record.id}"
    await client.put(url + "/review", json={"revision": 0, "issues": []})
    result = await client.post(url + "/export", json={"revision": 1, "format": "txt"})
    assert result.status_code == 200
    assert paths and not os.path.exists(os.path.dirname(paths[0]))


@pytest.mark.parametrize("mode", ["sync", "worker"])
async def test_document_full_original_saved_without_preaccepted_export(client, owner, tmp_path, monkeypatch, mode):
    from app.api.v1 import document as api
    from app.services import proofread as proofread_service
    from app.tasks.proofread_task import async_proofread_document
    text = "完整原文" * 25000
    document = await upload_for(owner, tmp_path, extracted_text=text)
    result = {"issues": [], "total_issues": 0, "usage": {}, "chunks_count": 1, "domain": "general",
              "coverage": {"status": "complete", "total_chunks": 1, "completed_chunks": 1, "failed_chunks": []},
              "depth": "standard", "config_id": 7}
    monkeypatch.setattr(api, "proofread_text", AsyncMock(return_value=result))
    monkeypatch.setattr(api, "charge_user_daily_quota", AsyncMock())
    monkeypatch.setattr(proofread_service, "proofread_text", AsyncMock(return_value=result))
    if mode == "sync":
        response = await client.post("/api/v1/document/proofread", json={"file_id": document.file_id})
        assert response.status_code == 200, response.text
        payload = response.json()
    else:
        async with async_session_factory() as db:
            task = ProofreadTask(task_id=str(uuid.uuid4()), document_id=document.file_id, owner_kind="user",
                                 owner_user_id=owner.id, params_json={"domain": "general"}, status="PENDING")
            db.add(task)
            await db.commit()
        payload = async_proofread_document.apply(args=(task.id,)).get()
    assert payload["corrected_download_url"] is None
    assert payload["coverage"] == result["coverage"] and payload["config_id"] == 7
    async with async_session_factory() as db:
        record = await db.get(ProofreadRecord, payload["record_id"])
        assert record.original_text == text and len(record.original_text) == 100000
        assert record.source_file_id == document.file_id


async def test_history_docx_uses_source_format_and_document_export_checks_live_owner(client, owner, tmp_path):
    path = tmp_path / "排版.docx"
    source = Document()
    paragraph = source.add_paragraph()
    paragraph.add_run("甲😀错词，乙").bold = True
    paragraph.add_run("错词").italic = True
    paragraph.add_run("。尾")
    source.save(path)
    document = await upload_for(owner, tmp_path, filename=path.name, file_ext=".docx", file_path=str(path))
    record = await record_for(owner, type="document", source_file_id=document.file_id, source_filename=path.name)
    url = f"/api/v1/history/{record.id}"
    await client.put(url + "/review", json={"revision": 0, "issues": [issue()]})
    exported = await client.post(url + "/export", json={"revision": 1, "format": "docx"})
    assert exported.status_code == 200, exported.text if exported.status_code != 200 else ""
    paragraph = Document(io.BytesIO(exported.content)).paragraphs[0]
    assert paragraph.text == "甲😀错词，乙正词。尾" and paragraph.runs[0].bold and paragraph.runs[1].italic
    doc_url = f"/api/v1/document/{document.file_id}/export"
    exported = await client.post(doc_url, json={"issues": [], "format": "docx"})
    assert exported.status_code == 200 and exported.content == path.read_bytes()
    # A warmed cache must not bypass current ownership/deletion state.
    async with async_session_factory() as db:
        row = await db.get(UploadedDocument, document.id)
        row.owner_kind = "apikey"
        await db.commit()
    assert (await client.post(doc_url, json={"issues": [], "format": "txt"})).status_code == 404
    async with async_session_factory() as db:
        row = await db.get(UploadedDocument, document.id)
        row.owner_kind = "user"
        row.status = "deleted"
        await db.commit()
    assert (await client.post(doc_url, json={"issues": [], "format": "txt"})).status_code == 404


async def test_export_failure_cleans_tempfile_and_hides_paths(client, owner, monkeypatch):
    import app.services.review as service
    paths = []

    def fail(*args):
        paths.append(args[3])
        raise ValueError(f"corrupt source at {args[3]}")

    monkeypatch.setattr(service, "_write_export", fail)
    record = await record_for(owner)
    url = f"/api/v1/history/{record.id}"
    await client.put(url + "/review", json={"revision": 0, "issues": []})
    result = await client.post(url + "/export", json={"revision": 1, "format": "docx"})
    assert result.status_code == 422 and "TXT" in result.text
    assert paths[0] not in result.text and not os.path.exists(os.path.dirname(paths[0]))


def compare_result():
    partial = {"status": "partial", "total_chunks": 2, "completed_chunks": 1,
               "failed_chunks": [{"chunk_index": 0, "start": 0, "end": 5, "text": SOURCE[:5],
                                  "error_code": "MODEL_ERROR"}]}
    models = [
        {"config_id": 11, "config_name": "模型甲", "model": "model-a", "domain": "legal", "depth": "deep",
         "success": True, "elapsed_ms": 123, "issues": [
             issue(explanation="甲的解释", source="llm", found_by=["模型甲"]),
             issue(start=None, end=None, chunk_index=0, _accepted=False, _ignored=True),
         ], "coverage": partial, "total_issues": 2, "complete": False},
        {"config_id": 22, "config_name": "模型乙", "model": "model-b", "domain": "official", "depth": "quick",
         "success": True, "elapsed_ms": 456, "issues": [
             issue(explanation="乙的独立解释", severity="error", _accepted=False),
             issue(type="grammar", explanation="乙的语法解释", _accepted=False),
             issue(start=None, end=None, chunk_index=9, _accepted=False),
         ], "coverage": {"status": "complete", "total_chunks": 1, "completed_chunks": 1, "failed_chunks": []}},
        {"config_id": 33, "config_name": "模型丙", "model": "model-c", "success": False,
         "error": "timeout", "elapsed_ms": 789, "issues": [],
         "coverage": {"status": "partial", "total_chunks": 1, "completed_chunks": 0,
                      "failed_chunks": [{"chunk_index": 0, "start": 0, "end": len(SOURCE), "text": SOURCE}]}},
    ]
    return {"compare": True, "results": models, "issues": models[0]["issues"],
            "domain": "legal", "depth": "deep", "consensus_originals": ["错词"], "only_in": {}}


def review_body(snapshot):
    return {key: copy.deepcopy(snapshot[key]) for key in ("revision", "issues", "coverage", "compare")}


async def test_compare_history_versions_retry_and_restore_without_model_or_billing(client, owner, monkeypatch):
    from app.services import model_compare, proofread

    no_model = AsyncMock(side_effect=AssertionError("review must not call LLM"))
    monkeypatch.setattr(proofread, "proofread_text", no_model)
    monkeypatch.setattr(model_compare, "run_proofread_compare", no_model)
    record = await record_for(owner, result=compare_result(), domain="legal", quota_weight=2)
    url = f"/api/v1/history/{record.id}"
    first_response = await client.get(url + "/review")
    assert first_response.status_code == 200, first_response.text
    first = first_response.json()
    a, b, failed = first["compare"]["results"]
    assert a["config_name"] == "模型甲" and a["model"] == "model-a"
    assert (a["domain"], a["depth"], a["elapsed_ms"]) == ("legal", "deep", 123)
    assert (b["domain"], b["depth"], b["elapsed_ms"]) == ("official", "quick", 456)
    assert a["issues"][0]["explanation"] == "甲的解释"
    assert b["issues"][0]["explanation"] == "乙的独立解释" and b["issues"][0]["severity"] == "error"
    assert a["coverage"]["status"] == "partial"
    assert failed["success"] is False and failed["error"] == "timeout"
    assert failed["domain"] == "general" and failed["depth"] == "standard"
    assert len(first["issues"]) == 3  # 未定位的相同问题忽略 chunk_index 后合并。
    for model in first["compare"]["results"]:
        assert set(model) == {"config_id", "config_name", "model", "domain", "depth", "success", "error",
                              "elapsed_ms", "issues", "coverage"}
        for item in model["issues"]:
            assert not item["_accepted"] and not item["_ignored"]
            assert "source" not in item and "found_by" not in item
    async with async_session_factory() as db:
        untouched = await db.get(ProofreadRecord, record.id)
        assert untouched.review_state is None and untouched.review_revision == 0

    body = review_body(first)
    for item in body["issues"]:
        item["_accepted"] = item["start"] is not None
        item["_ignored"] = item["start"] is None
    # 不同 type 的共享修改只执行一次；顶层解释无须与每模型解释相同。
    body["issues"][0]["explanation"] = "审阅备注"
    saved = await client.post(url + "/versions", json={**body, "label": "部分覆盖决策"})
    assert saved.status_code == 200, saved.text
    saved = saved.json()
    v1 = saved["versions"][0]
    assert saved["modified_text"] == "甲😀错词，乙正词。尾"
    assert v1["compare"] == first["compare"]
    assert (await client.get(url + "/review")).json() == saved
    assert (await client.put(url + "/review", json=body)).status_code == 409

    body = review_body(saved)
    # 模拟前端已付费补查并映回全文的结果；保存本身不作服务器覆盖认证。
    extra = issue(start=2, end=4, _accepted=False, explanation="补查报告")
    body["compare"]["results"][0]["issues"].append(extra)
    body["compare"]["results"][0]["coverage"] = {
        "status": "complete", "total_chunks": 2, "completed_chunks": 2, "failed_chunks": [],
    }
    body["issues"].append({**extra, "_accepted": True})
    saved = await client.post(url + "/versions", json=body)
    assert saved.status_code == 200, saved.text
    saved = saved.json()
    v2 = saved["versions"][1]
    assert saved["modified_text"] == "甲😀正词，乙正词。尾"
    assert saved["versions"][0] == v1
    assert v2["compare"]["results"][0]["coverage"]["status"] == "complete"
    assert (await client.get(url + "/review")).json() == saved
    for version in (v1, v2):
        restored = await client.put(url + "/review", json={
            "revision": saved["revision"],
            **{key: version[key] for key in ("issues", "coverage", "compare")},
        })
        assert restored.status_code == 200, restored.text
        saved = restored.json()
        assert saved["compare"] == version["compare"] and saved["issues"] == version["issues"]
        assert saved["modified_text"] == version["modified_text"]
        assert saved["versions"] == [v1, v2]
    # 省略 compare 沿用当前草稿（包括补查），仍保留版本。
    inherited = await client.put(url + "/review", json={"revision": saved["revision"], "issues": saved["issues"]})
    assert inherited.status_code == 200 and inherited.json()["compare"] == v2["compare"]
    async with async_session_factory() as db:
        stored = await db.get(ProofreadRecord, record.id)
        assert stored.result == record.result and stored.result["results"][0]["coverage"]["status"] == "partial"
        assert stored.quota_weight == 2 and stored.original_text == SOURCE
        assert await db.scalar(select(func.count()).select_from(ProofreadRecord).where(
            ProofreadRecord.user_id == owner.id)) == 1
    assert not await redis_module.redis_client.keys("textmirror:user_daily:*")
    assert not await redis_module.redis_client.keys("textmirror:apikey_daily:*")
    no_model.assert_not_called()


@pytest.mark.parametrize("invalid", [
    "model_set", "duplicate", "too_few", "too_many", "config_id", "config_name", "model", "domain", "depth",
    "elapsed_ms", "error", "success", "failed_success", "failed_issues", "failed_coverage",
    "nested_position", "nested_original", "failed_slice", "failed_index", "failed_count",
    "membership_missing", "membership_extra", "membership_type", "membership_location", "membership_suggestion",
    "nested_accepted", "nested_ignored", "shared_accepted", "shared_ignored", "shared_pending", "null",
    "omit_membership", "extra_state", "extra_model", "derived_total", "derived_complete", "too_many_issues",
])
async def test_compare_rejects_invalid_state_without_mutation(client, owner, invalid):
    record = await record_for(owner, result=compare_result())
    url = f"/api/v1/history/{record.id}/review"
    first = (await client.get(url)).json()
    body = review_body(first)
    models = body["compare"]["results"]
    if invalid in {"config_id", "config_name", "model", "domain", "depth", "elapsed_ms", "error", "success"}:
        models[0][invalid] = {"config_id": 0, "config_name": "伪造", "model": "forged", "domain": "general",
                              "depth": "standard", "elapsed_ms": 124, "error": "forged", "success": False}[invalid]
    elif invalid == "model_set":
        models[0]["config_id"] = 99
    elif invalid == "duplicate":
        models[1]["config_id"] = models[0]["config_id"]
    elif invalid == "too_few":
        body["compare"]["results"] = models[:1]
    elif invalid == "too_many":
        models.extend([{**models[0], "config_id": i} for i in (44, 55)])
    elif invalid == "failed_success":
        models[2]["success"] = True
    elif invalid == "failed_issues":
        models[2]["issues"].append(issue(_accepted=False))
    elif invalid == "failed_coverage":
        models[2]["coverage"] = None
    elif invalid == "nested_position":
        models[0]["issues"][0]["end"] = 999
    elif invalid == "nested_original":
        models[0]["issues"][0]["original"] = "伪造"
    elif invalid == "failed_slice":
        models[0]["coverage"]["failed_chunks"][0]["text"] = "伪造"
    elif invalid == "failed_index":
        models[0]["coverage"]["failed_chunks"][0]["chunk_index"] = 999
    elif invalid == "failed_count":
        models[0]["coverage"]["completed_chunks"] = 2
    elif invalid in {"membership_missing", "omit_membership"}:
        body["issues"] = []
        if invalid == "omit_membership":
            del body["compare"]
    elif invalid == "membership_extra":
        body["issues"].append(issue(start=2, end=4, _accepted=False))
    elif invalid == "membership_type":
        body["issues"][0]["type"] = "style"
    elif invalid == "membership_location":
        body["issues"][0].update(start=2, end=4)
    elif invalid == "membership_suggestion":
        body["issues"][0]["suggestion"] = "其他建议"
    elif invalid in {"nested_accepted", "nested_ignored"}:
        models[0]["issues"][0]["_accepted" if invalid == "nested_accepted" else "_ignored"] = True
    elif invalid.startswith("shared_"):
        body["issues"][0]["_accepted"] = True
        grammar = next(item for item in body["issues"] if item["type"] == "grammar")
        if invalid == "shared_ignored":
            grammar["_ignored"] = True
        elif invalid == "shared_accepted":
            body["issues"][0]["_accepted"] = False
            body["issues"][0]["_ignored"] = True
    elif invalid == "null":
        body["compare"] = None
    elif invalid == "extra_state":
        body["compare"]["consensus_originals"] = []
    elif invalid in {"extra_model", "derived_total", "derived_complete"}:
        models[0][{"extra_model": "only_in", "derived_total": "total_issues", "derived_complete": "complete"}[invalid]] = 1
    elif invalid == "too_many_issues":
        models[0]["issues"] = [models[0]["issues"][0]] * 5001
    response = await client.put(url, json=body)
    assert response.status_code == 422, response.text
    assert (await client.get(url)).json() == first
    async with async_session_factory() as db:
        stored = await db.get(ProofreadRecord, record.id)
        assert stored.review_state is None and stored.review_revision == 0 and stored.result == record.result


async def test_compare_nested_size_limit_and_inherited_state_limit(client, owner, monkeypatch):
    from app.services import review as review_service

    record = await record_for(owner, result=compare_result())
    url = f"/api/v1/history/{record.id}/review"
    first = (await client.get(url)).json()
    body = review_body(first)
    model = body["compare"]["results"][0]
    model["issues"] = [{**model["issues"][0], "explanation": "甲" * 20000}] * 150
    assert (await client.put(url, json=body)).status_code == 422
    # 省略 compare 不得绕过持久化草稿大小限制。
    monkeypatch.setattr(review_service, "MAX_REVIEW_BYTES", 1000)
    assert (await client.put(url, json={"revision": 0, "issues": first["issues"]})).status_code == 422
    assert (await client.get(url)).json() == first


@pytest.mark.parametrize("legacy", [False, True])
async def test_single_or_legacy_compare_records_do_not_invent_models(client, owner, legacy):
    result = {"issues": [issue(_accepted=False)]}
    if legacy:
        result["compare"] = True
        result["models"] = ["老模型甲", "老模型乙"]
    record = await record_for(owner, result=result)
    url = f"/api/v1/history/{record.id}"
    first = (await client.get(url + "/review")).json()
    assert first["compare"] is None
    forged = {"results": [{key: value for key, value in model.items() if key not in {"total_issues", "complete"}}
                           for model in compare_result()["results"]]}
    for model in forged["results"]:
        model["issues"] = []
    assert (await client.put(url + "/review", json={"revision": 0, "issues": [], "compare": forged})).status_code == 422
    saved = await client.post(url + "/versions", json={"revision": 0, "issues": [issue()]})
    assert saved.status_code == 200 and saved.json()["compare"] is None
    assert saved.json()["versions"][0]["compare"] is None


@pytest.mark.parametrize("accepted,ignored", [(True, False), (False, True), (False, False)])
async def test_existing_generic_compare_draft_and_legacy_version_recover_models(client, owner, accepted, ignored):
    state = {"issues": [issue(_accepted=accepted, _ignored=ignored)], "coverage": None, "versions": [{
        "id": "legacy", "label": "旧版本", "created_at": "2026-01-01T00:00:00Z", "issues": [], "modified_text": SOURCE,
    }]}
    record = await record_for(owner, result=compare_result(), review_state=state, review_revision=1)
    url = f"/api/v1/history/{record.id}/review"
    first = (await client.get(url)).json()
    assert len(first["compare"]["results"]) == 3
    assert first["issues"][0]["_accepted"] is accepted
    assert first["versions"][0]["compare"] is None
    assert len(first["issues"]) == 3
    located = [item for item in first["issues"] if item["start"] is not None]
    assert {item["type"] for item in located} == {"typo", "grammar"}
    assert all((item["_accepted"], item["_ignored"]) == (accepted, ignored) for item in located)
    assert all(not item["_accepted"] and not item["_ignored"] for model in first["compare"]["results"]
               for item in model["issues"])
    async with async_session_factory() as db:
        assert (await db.get(ProofreadRecord, record.id)).review_state == state
    saved = await client.put(url, json=review_body(first))
    assert saved.status_code == 200, saved.text
    assert saved.json()["issues"] == first["issues"]
    assert saved.json()["compare"] == first["compare"]
    assert saved.json()["versions"] == first["versions"]


async def test_compare_cas_rejects_stale_report_and_keeps_version_ceiling(client, owner):
    record = await record_for(owner, result=compare_result())
    url = f"/api/v1/history/{record.id}"
    first = (await client.get(url + "/review")).json()
    body = review_body(first)
    async with async_session_factory() as reader, async_session_factory() as writer:
        stale = await load_review_record(reader, record.id, owner.id)
        await reader.commit()
        latest = await load_review_record(writer, record.id, owner.id)
        await save_review(writer, latest, owner.id, ReviewWriteRequest.model_validate(body), create_version=True)
        await writer.commit()
        with pytest.raises(HTTPException) as failure:
            await save_review(reader, stale, owner.id, ReviewWriteRequest.model_validate(body))
        assert failure.value.status_code == 409
        await reader.rollback()
    for revision in range(1, 20):
        body["revision"] = revision
        response = await client.post(url + "/versions", json=body)
        assert response.status_code == 200, response.text
    full = response.json()
    assert len(full["versions"]) == 20
    assert all(version["compare"] == first["compare"] for version in full["versions"])
    body["revision"] = 20
    assert (await client.post(url + "/versions", json=body)).status_code == 422
    assert (await client.get(url + "/review")).json() == full
    assert (await client.put(url + "/review", json=body)).status_code == 200


async def history_pair(client, record):
    listing = await client.get("/api/v1/history", params={"page_size": 100})
    assert listing.status_code == 200, listing.text
    item = next(item for item in listing.json()["items"] if item["id"] == record.id)
    detail = await client.get(f"/api/v1/history/{record.id}")
    assert detail.status_code == 200, detail.text
    detail = detail.json()
    for key in ("mode", "coverage_status", "review_summary"):
        assert item[key] == detail[key]
    return item, detail


@pytest.mark.parametrize("kind", ["text", "document"])
async def test_history_initial_and_saved_decisions_share_review_semantics(client, owner, kind):
    record = await record_for(owner, type=kind, total_issues=99, result={"issues": [issue()]})
    item, detail = await history_pair(client, record)
    assert item["mode"] == "single" and item["coverage_status"] == "unknown"
    assert item["review_summary"] == {"total": 1, "accepted": 0, "ignored": 0, "pending": 1, "failed_models": 0}
    assert detail["issues"][0]["_accepted"] is False  # 模型 flags 不是用户决策
    url = f"/api/v1/history/{record.id}/review"
    saved = await client.put(url, json={"revision": 0, "issues": [
        issue(), issue(start=2, end=4, _accepted=False, _ignored=True),
        issue(start=None, end=None, _accepted=False),
    ]})
    assert saved.status_code == 200, saved.text
    item, detail = await history_pair(client, record)
    assert item["review_summary"] == {"total": 3, "accepted": 1, "ignored": 1, "pending": 1, "failed_models": 0}
    assert detail["issues"] == saved.json()["issues"]
    assert item["total_issues"] == 99  # 旧契约保留，UI 使用当前草稿摘要


@pytest.mark.parametrize("legacy_draft", [False, True])
async def test_history_compare_merges_models_and_inherits_legacy_decisions(client, owner, legacy_draft):
    result = compare_result()
    del result["issues"]
    # 失败模型即使携带建议也不能混入已发现问题。
    result["results"][2]["issues"] = [issue(type="failed-only")]
    state = {"issues": [issue()], "coverage": None} if legacy_draft else None
    record = await record_for(owner, result=result, review_state=state)
    item, detail = await history_pair(client, record)
    assert item["mode"] == "compare" and item["coverage_status"] == "partial"
    assert item["review_summary"] == {
        "total": 3, "accepted": 2 if legacy_draft else 0, "ignored": 0,
        "pending": 1 if legacy_draft else 3, "failed_models": 1,
    }
    assert {issue["type"] for issue in detail["issues"]} == {"typo", "grammar"}
    restored = (await client.get(f"/api/v1/history/{record.id}/review")).json()
    assert detail["issues"] == restored["issues"]


@pytest.mark.parametrize("coverage_status", ["complete", "partial", "unknown"])
async def test_history_compare_coverage_comes_from_each_saved_model(client, owner, coverage_status):
    result = compare_result()
    result["results"] = result["results"][:2]
    record = await record_for(owner, result=result)
    url = f"/api/v1/history/{record.id}/review"
    body = review_body((await client.get(url)).json())
    for model in body["compare"]["results"]:
        model["coverage"] = {"status": "complete", "total_chunks": 1, "completed_chunks": 1, "failed_chunks": []}
    if coverage_status == "unknown":
        body["compare"]["results"][0]["coverage"] = None
    elif coverage_status == "partial":
        body["compare"]["results"][0]["coverage"] = result["results"][0]["coverage"]
    assert (await client.put(url, json=body)).status_code == 200
    item, _ = await history_pair(client, record)
    assert item["coverage_status"] == coverage_status
    assert item["review_summary"]["failed_models"] == 0


@pytest.mark.parametrize("status,coverage,expected", [
    ("complete", "complete", "complete"), ("partial", "complete", "partial"),
    ("complete", None, "unknown"), ("complete", "partial", "partial"),
])
async def test_history_collaboration_uses_merged_findings_not_role_counts(client, owner, status, coverage, expected):
    result = {"collaboration": {"status": status, "roles": [
        {"id": "language", "status": "success", "issue_count": 8},
        {"id": "consistency", "status": "success", "issue_count": 8},
    ], "findings": [issue(found_by=["language", "consistency"], review_status="confirmed")]}}
    if coverage:
        result["coverage"] = {"status": coverage}
    record = await record_for(owner, result=result)
    item, detail = await history_pair(client, record)
    assert item["mode"] == "collaboration" and item["coverage_status"] == expected
    assert item["review_summary"]["total"] == item["review_summary"]["pending"] == 1
    assert not detail["issues"][0]["_accepted"]
    # 已保存空草稿不回填原始 findings。
    async with async_session_factory() as db:
        stored = await db.get(ProofreadRecord, record.id)
        stored.review_state = {"issues": [], "coverage": None}
        await db.commit()
    item, detail = await history_pair(client, record)
    assert item["review_summary"]["total"] == 0 and detail["issues"] == []


async def test_history_legacy_compare_and_failed_models_never_claim_complete(client, owner):
    legacy = await record_for(owner, result={"compare": True, "models": ["甲", "乙"], "issues": []})
    item, _ = await history_pair(client, legacy)
    assert item["mode"] == "compare" and item["coverage_status"] == "unknown"
    result = compare_result()
    for model in result["results"]:
        model.update(success=False, issues=[], coverage=None)
    failed = await record_for(owner, result=result)
    item, detail = await history_pair(client, failed)
    assert item["coverage_status"] == "partial" and item["review_summary"]["failed_models"] == 3
    assert item["review_summary"]["total"] == 0 and detail["issues"] == []


async def test_history_list_projects_small_fields_in_two_queries_and_enforces_owner(client, owner, monkeypatch):
    from sqlalchemy import event

    from app.api.v1 import history as api
    from app.core.database import engine
    from app.services import review as service

    def forbidden(*args, **kwargs):
        raise AssertionError("history must not build a full ReviewResponse")

    monkeypatch.setattr(service, "review_response", forbidden)
    monkeypatch.setattr(api, "review_response", forbidden)
    for _ in range(3):
        await record_for(owner, original_text="甲" * 100000, result={"issues": []},
                         review_state={"issues": [], "versions": [{"invalid": "巨型版本" * 100000}]})
    foreign = await record_for(owner, user_id=None)
    polish = await record_for(owner, type="polish", result={"versions": [{"label": "润色", "content": "正文"}]})
    queries = []

    def capture(conn, cursor, statement, parameters, context, executemany):
        queries.append((statement, parameters))

    event.listen(engine.sync_engine, "before_cursor_execute", capture)
    try:
        response = await client.get("/api/v1/history", params={"type": "text", "domain": "general", "page_size": 2})
    finally:
        event.remove(engine.sync_engine, "before_cursor_execute", capture)
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total"] == 3 and len(payload["items"]) == 2
    assert len(queries) == 2
    sql, params = queries[1]
    assert "substr(proofread_records.original_text" in sql
    assert "proofread_records.modified_text" not in sql
    assert ", proofread_records.review_state" not in sql and "versions" not in str(params)
    assert all(len(item["text_preview"]) == 103 for item in payload["items"])
    assert len(response.content) < 3000 and "巨型版本" not in response.text
    second = (await client.get("/api/v1/history", params={"type": "text", "page_size": 2, "page": 2})).json()
    assert len(second["items"]) == 1
    assert {item["id"] for item in payload["items"]}.isdisjoint(item["id"] for item in second["items"])
    assert (await client.get(f"/api/v1/history/{foreign.id}")).status_code == 404
    item, detail = await history_pair(client, polish)
    assert item["mode"] is None and detail["result"] == polish.result
