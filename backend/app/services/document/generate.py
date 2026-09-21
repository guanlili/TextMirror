"""按校对结果生成修订版文档（DOCX / TXT）。"""
import shutil
from typing import List

from docx import Document as DocxDocument
from docx.shared import RGBColor
from loguru import logger


class ReviewDocxError(ValueError):
    """可向用户展示的安全导出错误（不包含源文件路径）。"""


def generate_reviewed_docx(
    original_path: str,
    source_text: str,
    patches: List[dict],
    output_path: str,
) -> str:
    """按 _extract_docx 的 trim/非空段落语义定位，仅修改命中的 XML 文本片段。

    不清空 run、不改变样式/链接/段落。跨段或控制字符替换无法保证视觉与
    TXT 成稿语义一致，明确拒绝；调用方可改为导出 TXT。
    """
    from docx.oxml.ns import qn

    error = "该修改无法在保留 DOCX 格式时保持文本一致（跨段、换行或特殊内容）；请导出 TXT"
    doc = DocxDocument(original_path)
    paragraphs = [(p, p.text) for p in doc.paragraphs if p.text.strip()]
    if "\n".join(text.strip() for _, text in paragraphs) != source_text:
        raise ReviewDocxError("源 DOCX 文本与审阅原文不一致；请导出 TXT")

    patches = sorted(patches, key=lambda patch: patch["start"])
    cursor = 0
    parts = []
    for patch in patches:
        start, end = patch["start"], patch["end"]
        if not (0 <= start < end <= len(source_text)) or start < cursor or source_text[start:end] != patch["original"]:
            raise ReviewDocxError("审阅补丁定位无效或重叠")
        parts.extend((source_text[cursor:start], patch["replacement"]))
        cursor = end
    parts.append(source_text[cursor:])
    expected_text = "".join(parts)
    if not patches:
        shutil.copyfile(original_path, output_path)
        return output_path

    # 全局码点 -> 段落原始文本 -> w:t 内偏移。保留被 strip 的空白与空段落。
    spans = []
    position = 0
    for paragraph, raw in paragraphs:
        trimmed = raw.strip()
        spans.append((position, position + len(trimmed), paragraph, raw, len(raw) - len(raw.lstrip())))
        position += len(trimmed) + 1

    # Word 会按域指令重算缓存文本，不能把缓存改动当作稳定成稿。
    field_text_nodes = set()
    field_depth = 0
    for node in doc.element.body.iter():
        if node.tag == qn("w:fldChar"):
            field_type = node.get(qn("w:fldCharType"))
            if field_type == "begin":
                field_depth += 1
            elif field_type == "end":
                field_depth = max(0, field_depth - 1)
        elif node.tag == qn("w:t") and field_depth:
            field_text_nodes.add(node)

    def text_nodes(paragraph, raw):
        def collect(runs):
            nodes = []
            offset = 0
            for run in runs:
                if run.tag == qn("w:fldSimple"):
                    nodes.append((offset, offset, run, ""))
                    continue
                for child in run:
                    if child.tag == qn("w:t"):
                        text = child.text or ""
                    elif child.tag == qn("w:tab"):
                        text = "\t"
                    elif child.tag in (qn("w:br"), qn("w:cr")):
                        text = "\n"
                    elif child.tag in (qn("w:fldChar"), qn("w:instrText")):
                        text = ""
                    else:
                        continue
                    nodes.append((offset, offset + len(text), child, text))
                    offset += len(text)
            return nodes

        nodes = collect(paragraph._p.xpath("./w:r | ./w:hyperlink/w:r | ./w:fldSimple | ./w:hyperlink/w:fldSimple"))
        if "".join(item[3] for item in nodes) != raw:
            # Older python-docx excludes hyperlink text in Paragraph.text; honor that range.
            nodes = collect(paragraph._p.xpath("./w:r | ./w:fldSimple"))
        if "".join(item[3] for item in nodes) != raw:
            raise ReviewDocxError(error)
        return nodes

    # 缓存原始节点坐标；逆序编辑确保较早位置不受后续文本长度变化影响。
    mapped = {}
    for patch in reversed(patches):
        if patch["replacement"] == patch["original"]:
            continue
        span = next((span for span in spans if span[0] <= patch["start"] < patch["end"] <= span[1]), None)
        if span is None or any(char in patch["replacement"] for char in "\r\n\t"):
            raise ReviewDocxError(error)
        begin, _, paragraph, raw, trim = span
        if begin not in mapped:
            mapped[begin] = text_nodes(paragraph, raw)
        start, end = patch["start"] - begin + trim, patch["end"] - begin + trim
        touched = [node for node in mapped[begin] if node[0] < end and node[1] > start]
        if not touched or any(node[2].tag != qn("w:t") or node[2] in field_text_nodes for node in touched):
            raise ReviewDocxError(error)
        for index, (left, right, node, _) in enumerate(touched):
            old = node.text or ""
            a, b = max(start - left, 0), min(end - left, right - left)
            node.text = old[:a] + (patch["replacement"] if index == 0 else "") + old[b:]
            node.set(qn("xml:space"), "preserve")

    # strip/空段过滤可能使删除后语义改变；此时绝不交付与 TXT 不一致的 DOCX。
    actual = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    if actual != expected_text:
        raise ReviewDocxError(error)
    doc.save(output_path)
    return output_path


def generate_corrected_docx(
    original_path: str,
    issues: List[dict],
    output_path: str,
) -> str:
    """
    基于校对结果生成修订版 Word 文档
    在原文基础上用红色标记修改内容

    :param original_path: 原始 docx 文件路径
    :param issues: 校对问题列表
    :param output_path: 输出文件路径
    :return: 输出文件路径
    """
    doc = DocxDocument(original_path)

    # 构建替换映射
    replacements = {}
    for issue in issues:
        original = issue.get("original", "")
        suggestion = issue.get("suggestion", "")
        if original and suggestion and original != suggestion:
            replacements[original] = suggestion

    # 遍历段落进行替换
    for para in doc.paragraphs:
        full_text = para.text
        modified = False
        for original, suggestion in replacements.items():
            if original in full_text:
                full_text = full_text.replace(original, suggestion)
                modified = True

        if modified:
            # 清除原有 run，用新文本重建
            for run in para.runs:
                run.text = ""
            if para.runs:
                para.runs[0].text = full_text
                para.runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
            else:
                run = para.add_run(full_text)
                run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    doc.save(output_path)
    logger.info(f"修订文档已生成: {output_path}")
    return output_path


def generate_corrected_txt(
    original_text: str,
    issues: List[dict],
    output_path: str,
) -> str:
    """
    基于校对结果生成修订版 TXT 文件

    :param original_text: 原始文本
    :param issues: 校对问题列表
    :param output_path: 输出文件路径
    :return: 输出文件路径
    """
    corrected = original_text
    for issue in issues:
        original = issue.get("original", "")
        suggestion = issue.get("suggestion", "")
        if original and suggestion and original != suggestion:
            corrected = corrected.replace(original, suggestion)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(corrected)

    logger.info(f"修订文本已生成: {output_path}")
    return output_path
