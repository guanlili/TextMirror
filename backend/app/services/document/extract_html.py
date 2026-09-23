"""从 DOCX 文件中提取格式化 HTML，保留排版和字体样式。"""
from html import escape as html_escape

from docx import Document as DocxDocument
from loguru import logger

from .css_sanitize import _safe_css_color, _safe_css_font_family, _safe_css_length


def extract_html_from_file(file_path: str, file_ext: str, plain_text: str = "") -> str:
    """
    从文件中提取格式化 HTML，保留排版和字体样式
    Word 文档完整保留样式，TXT/PDF 使用纯文本包装

    :param file_path: 文件路径
    :param file_ext: 文件扩展名（含点，如 .docx）
    :param plain_text: 已提取的纯文本（用于 txt/pdf 的 HTML 包装）
    :return: HTML 字符串
    """
    ext = file_ext.lower()
    if ext == ".docx":
        return _extract_docx_html(file_path)
    elif ext == ".doc":
        # .doc 格式无法直接提取富文本样式，降级为纯文本 HTML 包装
        escaped = html_escape(plain_text)
        return f'<div style="white-space:pre-wrap;line-height:1.8;font-size:14px;">{escaped}</div>'
    else:
        escaped = html_escape(plain_text)
        return f'<div style="white-space:pre-wrap;line-height:1.8;font-size:14px;">{escaped}</div>'


def _extract_docx_html(file_path: str) -> str:
    """
    将 Word 文档转换为 HTML，保留排版和字体样式
    支持：段落样式、标题层级、对齐方式、首行缩进、段间距、行间距、
          加粗、斜体、下划线、删除线、字号、字体、颜色、表格
    """
    doc = DocxDocument(file_path)
    html_parts = []

    # 按文档顺序处理段落和表格
    para_idx = 0
    table_idx = 0

    for child in doc.element.body:
        tag_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag

        if tag_name == 'p':
            if para_idx < len(doc.paragraphs):
                html_parts.append(_para_to_html(doc.paragraphs[para_idx]))
                para_idx += 1
        elif tag_name == 'tbl':
            if table_idx < len(doc.tables):
                html_parts.append(_table_to_html(doc.tables[table_idx]))
                table_idx += 1

    return '\n'.join(html_parts)


def _para_to_html(para) -> str:
    """将段落转换为 HTML 标签，保留排版样式"""
    # 空段落保留为空行
    if not para.text.strip():
        return '<p style="margin:0.3em 0;min-height:1em;"><br/></p>'

    # 根据样式确定标签（标题 → h1-h6，正文 → p）
    tag = 'p'
    try:
        style_name = (para.style.name or '') if para.style else ''
        for i in range(1, 7):
            if f'Heading {i}' in style_name or style_name == f'Heading{i}':
                tag = f'h{i}'
                break
    except Exception as e:
        logger.debug(f"[docx html] 标题检测异常: {e}")

    # 收集段落样式
    styles = ['margin:0.3em 0']

    # 对齐方式
    try:
        if para.alignment is not None:
            align_val = int(para.alignment)
            align_map = {1: 'center', 2: 'right', 3: 'justify', 4: 'justify', 5: 'justify'}
            align = align_map.get(align_val, '')
            if align:
                styles.append(f'text-align:{align}')
    except Exception as e:
        logger.debug(f"[docx html] 对齐方式读取异常: {e}")

    # 首行缩进
    try:
        pf = para.paragraph_format
        if pf.first_line_indent and pf.first_line_indent.pt > 0:
            decl = _safe_css_length('text-indent', pf.first_line_indent.pt)
            if decl:
                styles.append(decl)
    except Exception as e:
        logger.debug(f"[docx html] 首行缩进读取异常: {e}")

    # 左缩进
    try:
        pf = para.paragraph_format
        if pf.left_indent and pf.left_indent.pt > 0:
            decl = _safe_css_length('padding-left', pf.left_indent.pt)
            if decl:
                styles.append(decl)
    except Exception as e:
        logger.debug(f"[docx html] 左缩进读取异常: {e}")

    # 段前段后间距
    try:
        pf = para.paragraph_format
        if pf.space_before and pf.space_before.pt:
            decl = _safe_css_length('margin-top', pf.space_before.pt)
            if decl:
                styles.append(decl)
        if pf.space_after and pf.space_after.pt:
            decl = _safe_css_length('margin-bottom', pf.space_after.pt)
            if decl:
                styles.append(decl)
    except Exception as e:
        logger.debug(f"[docx html] 段间距读取异常: {e}")

    # 行间距
    try:
        pf = para.paragraph_format
        if pf.line_spacing is not None:
            if isinstance(pf.line_spacing, (int, float)):
                spacing = float(pf.line_spacing)
                if 0 < spacing <= 10:
                    styles.append(f'line-height:{spacing:.2f}')
            elif hasattr(pf.line_spacing, 'pt') and pf.line_spacing.pt:
                decl = _safe_css_length('line-height', pf.line_spacing.pt)
                if decl:
                    styles.append(decl)
    except Exception as e:
        logger.debug(f"[docx html] 行间距读取异常: {e}")

    # 构建行内 HTML（保留字体样式）
    inline_html = _runs_to_html(para.runs)
    if not inline_html:
        inline_html = html_escape(para.text)

    style_attr = f' style="{html_escape(";".join(styles), quote=True)}"' if styles else ''
    return f'<{tag}{style_attr}>{inline_html}</{tag}>'


def _runs_to_html(runs) -> str:
    """将文档 Runs 转换为行内 HTML，保留字体样式"""
    parts = []
    for run in runs:
        text = run.text
        if not text:
            continue

        # HTML 转义
        text = html_escape(text)

        # 收集行内样式（值一律走白名单，防止逃逸 style 属性注入事件处理器）
        run_styles = []

        # 字体名称
        try:
            decl = _safe_css_font_family(run.font.name)
            if decl:
                run_styles.append(decl)
        except Exception as e:
            logger.debug(f"[docx html] 字体名称读取异常: {e}")

        # 字号
        try:
            if run.font.size and run.font.size.pt:
                decl = _safe_css_length("font-size", run.font.size.pt, max_pt=200.0)
                if decl:
                    run_styles.append(decl)
        except Exception as e:
            logger.debug(f"[docx html] 字号读取异常: {e}")

        # 字体颜色
        try:
            if run.font.color and run.font.color.rgb:
                decl = _safe_css_color(run.font.color.rgb)
                if decl:
                    run_styles.append(decl)
        except Exception as e:
            logger.debug(f"[docx html] 字体颜色读取异常: {e}")

        # 加粗
        try:
            if run.bold:
                text = f'<strong>{text}</strong>'
        except Exception as e:
            logger.debug(f"[docx html] 加粗检测异常: {e}")

        # 斜体
        try:
            if run.italic:
                text = f'<em>{text}</em>'
        except Exception as e:
            logger.debug(f"[docx html] 斜体检测异常: {e}")

        # 下划线
        try:
            if run.underline:
                text = f'<u>{text}</u>'
        except Exception as e:
            logger.debug(f"[docx html] 下划线检测异常: {e}")

        # 删除线
        try:
            if run.font.strike:
                text = f'<s>{text}</s>'
        except Exception as e:
            logger.debug(f"[docx html] 删除线检测异常: {e}")

        # 包裹行内样式
        if run_styles:
            style_str = html_escape(';'.join(run_styles), quote=True)
            text = f'<span style="{style_str}">{text}</span>'

        parts.append(text)

    return ''.join(parts)


def _table_to_html(table) -> str:
    """将表格转换为 HTML"""
    html = '<table style="border-collapse:collapse;width:100%;margin:8px 0;">'
    for row in table.rows:
        html += '<tr>'
        for cell in row.cells:
            cell_text = html_escape(cell.text.strip())
            html += f'<td style="border:1px solid #ccc;padding:6px 8px;vertical-align:top;">{cell_text}</td>'
        html += '</tr>'
    html += '</table>'
    return html
